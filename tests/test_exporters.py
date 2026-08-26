"""Group 4a — exporter unit tests (backend.exporters.export_posts).

These tests call the exporters directly (posts_loader-based), so they are
independent of the API layer / DB: they verify file-level correctness —
JSON nested round-trip, CSV BOM + column set + escaping, XLSX sheet
structure/freeze panes/auto-filter/date formatting — exactly per spec §5.
"""
from __future__ import annotations

import csv
import io
import json
from datetime import datetime

from backend.exporters import SUPPORTED_FORMATS, export_posts
from backend.exporters.csv_exporter import FLAT_COLUMNS

from helpers import canonical_post, post_with_text, sample_posts

#: All 33 canonical keys must survive the JSON export untouched.
EXPECTED_KEYS = [
    "post_id", "facebook_url", "post_url", "page_name", "page_id",
    "profile_url", "post_type", "published_at", "timestamp", "text",
    "caption", "hashtags", "mentions", "external_links", "likes", "reactions",
    "comments_count", "shares", "views_count", "reaction_like_count",
    "reaction_love_count", "reaction_care_count", "reaction_haha_count",
    "reaction_wow_count", "reaction_sad_count", "reaction_angry_count",
    "media_type", "thumbnail_url", "media_url", "video_url", "transcript",
    "transcript_language", "scraped_at",
]


def test_json_nested_roundtrip(tmp_path):
    posts = sample_posts(3)
    path = export_posts(
        posts_loader=lambda: posts, job_id="j1", fmt="json", base_dir=tmp_path
    )
    assert path.is_file()
    with open(path, encoding="utf-8") as fh:
        loaded = json.load(fh)
    assert loaded == posts, "JSON export must be a lossless round-trip"
    assert all(list(p.keys()) == EXPECTED_KEYS for p in loaded)


def test_csv_bom_columns_and_escaping(tmp_path):
    tricky = post_with_text(
        'Line 1\nLine 2, with "quotes" and semicolon; and more'
    )
    posts = [tricky, canonical_post("1002", post_type="image")]
    path = export_posts(
        posts_loader=lambda: posts, job_id="j1", fmt="csv", base_dir=tmp_path
    )
    raw = path.read_bytes()
    assert raw[:3] == b"\xef\xbb\xbf", "CSV must start with a UTF-8 BOM"

    decoded = raw.decode("utf-8-sig")
    reader = list(csv.reader(io.StringIO(decoded)))
    assert reader[0] == FLAT_COLUMNS, "header must equal FLAT_COLUMNS"
    assert len(reader[0]) == 28
    assert len(reader) == 1 + len(posts)

    header = reader[0]
    row = reader[1]
    text_idx = header.index("text")
    assert row[text_idx] == tricky["text"], "multiline/quote/comma must survive"
    assert row[header.index("hashtags")] == "#a|#b"
    assert row[header.index("comments_count")] == "3"
    # None -> empty cell
    assert row[header.index("views_count")] == ""
    assert row[header.index("transcript")] == ""


def test_xlsx_structure(tmp_path):
    posts = sample_posts(3)
    path = export_posts(
        posts_loader=lambda: posts, job_id="j1", fmt="excel", base_dir=tmp_path
    )
    assert path.is_file() and path.suffix == ".xlsx"

    from openpyxl import load_workbook

    wb = load_workbook(path)
    assert wb.sheetnames == ["Posts", "Engagement", "Media", "Metadata"]

    ws_posts = wb["Posts"]
    assert ws_posts.freeze_panes == "A2"
    assert ws_posts.auto_filter.ref, "auto_filter must be set"
    header = [c.value for c in ws_posts[1]]
    assert header == FLAT_COLUMNS
    assert len(header) == 28

    posts_header = {name: i + 1 for i, name in enumerate(header)}
    # published_at formatted as a (tz-stripped) datetime with the date format
    cell = ws_posts.cell(row=2, column=posts_header["published_at"])
    assert isinstance(cell.value, datetime)
    assert cell.number_format == "yyyy-mm-dd hh:mm:ss"
    # URL columns become hyperlinks
    url_cell = ws_posts.cell(row=2, column=posts_header["post_url"])
    assert url_cell.hyperlink is not None
    assert url_cell.hyperlink.target == posts[0]["post_url"]

    ws_eng = wb["Engagement"]
    eng_header = [c.value for c in ws_eng[1]]
    assert "post_id" in eng_header and "comments_count" in eng_header
    assert set(eng_header) >= {
        "likes", "reactions", "shares", "views_count",
        "reaction_like_count", "reaction_angry_count",
    }

    ws_media = wb["Media"]
    media_header = [c.value for c in ws_media[1]]
    assert media_header == [
        "post_id", "media_type", "thumbnail_url", "media_url", "video_url",
    ]

    ws_meta = wb["Metadata"]
    meta = {r[0].value: r[1].value for r in ws_meta.iter_rows(min_row=2) if r[0].value}
    assert meta["posts_count"] == 3
    assert meta["job_id"] == "j1"
    assert meta["sheets"] == "Posts, Engagement, Media, Metadata"


def test_xlsx_via_xlsx_alias(tmp_path):
    posts = sample_posts(1)
    path = export_posts(
        posts_loader=lambda: posts, job_id="j1", fmt="xlsx", base_dir=tmp_path
    )
    assert path.name == "facebook_posts.xlsx"


def test_unsupported_format_rejected(tmp_path):
    import pytest

    with pytest.raises(ValueError, match="unsupported export format"):
        export_posts(posts_loader=lambda: [], fmt="html", base_dir=tmp_path)


def test_missing_posts_source_rejected(tmp_path):
    import pytest

    with pytest.raises(ValueError, match="posts_loader or posts_path"):
        export_posts(fmt="json", base_dir=tmp_path)


def test_non_dict_post_rejected(tmp_path):
    import pytest

    with pytest.raises(TypeError):
        export_posts(posts_loader=lambda: [{"ok": 1}, "nope"], fmt="json",
                     base_dir=tmp_path)


def test_supported_formats_exposed():
    assert SUPPORTED_FORMATS == ("json", "csv", "excel", "xlsx", "jsonl")